"""Regenerates the small preview fixtures next to this file (the outputs are committed).

    python3 make-fixtures.py        # needs python-docx, python-pptx, openpyxl
"""
import pathlib

import docx
import openpyxl
import pptx
from pptx.util import Inches

here = pathlib.Path(__file__).parent


def pdf(path: pathlib.Path, pages: list[str]) -> None:
    objs: list[bytes] = []
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(pages):
        content = f"BT /F1 28 Tf 72 700 Td ({text}) Tj ET".encode()
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R /Resources << /Font << /F1 3 0 R >> >> >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    out = b"%PDF-1.4\n"
    offsets = []
    for n, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % n + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(out)


pdf(here / "sample.pdf", ["Orbit PDF page one", "Orbit PDF page two"])

doc = docx.Document()
doc.add_heading("Orbit DOCX heading", 1)
doc.add_paragraph("Orbit DOCX paragraph about the release plan.")
doc.save(here / "sample.docx")

book = openpyxl.Workbook()
first = book.active
first.title = "Budget"
first.append(["Item", "Cost"])
first.append(["Orbit license", 1200])
first.append(["Hosting", 340])
second = book.create_sheet("Notes")
second.append(["Owner", "Orbit sheet two"])
book.save(here / "sample.xlsx")

(here / "sample.csv").write_text("name,city\nOrbit CSV row,Shanghai\nSecond row,Beijing\n", encoding="utf-8")

deck = pptx.Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = "Orbit PPTX title"
box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(6), Inches(1))
box.text_frame.text = "Orbit PPTX body text"
deck.save(here / "sample.pptx")

(here / "main.py").write_text('def greet(name: str) -> str:\n    """Say hello."""\n    return f"hello {name}"\n\n\nprint(greet("orbit"))\n', encoding="utf-8")
